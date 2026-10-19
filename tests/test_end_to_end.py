"""Full pipeline: generate files -> load -> reconcile -> report -> score."""
from openpyxl import load_workbook

from recon.cli import ROOT, main, reconcile
from recon.evaluate import evaluate
from recon.generate import generate
from recon.loader import load_records
from recon.config import ReconConfig


def test_loader_rejects_bad_rows_with_reasons(tmp_path):
    generate(tmp_path, n=200, seed=5)
    cfg = ReconConfig.load(ROOT / "config.json")
    _, rejected = load_records(tmp_path / "processor_settlement.csv", "processor", cfg.processor)
    reasons = " | ".join(r.reason for r in rejected)
    assert len(rejected) == 3
    assert "does not match format" in reasons and "amount" in reasons


def test_accuracy_on_generated_data(tmp_path):
    generate(tmp_path, n=1500, seed=21)
    res = reconcile(tmp_path / "internal_ledger.csv", tmp_path / "processor_settlement.csv", ROOT / "config.json")
    ev = evaluate(res, tmp_path / "ground_truth.json")
    assert ev["link_precision"] >= 0.99
    assert ev["link_recall"] >= 0.99
    assert ev["status_accuracy"] >= 0.99


def test_cli_writes_report(tmp_path):
    data, out = tmp_path / "data", tmp_path / "out"
    assert main(["generate", "--out", str(data), "--n", "300", "--seed", "8"]) == 0
    code = main(["run", "--internal", str(data / "internal_ledger.csv"),
                 "--processor", str(data / "processor_settlement.csv"), "--out", str(out)])
    assert code == 0
    wb = load_workbook(out / "recon_report.xlsx")
    assert wb.sheetnames == ["Summary", "Matched", "Breaks", "Unmatched_Internal",
                             "Unmatched_Processor", "Duplicates", "Rejected"]
    assert (out / "recon_results.csv").stat().st_size > 0
