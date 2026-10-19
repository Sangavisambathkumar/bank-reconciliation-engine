"""Write the reconciliation results to an Excel workbook and a flat CSV.

Sheets:
  Summary              headline numbers + the reconciliation proof (ties to zero)
  Matched              every matched group, one row per record
  Breaks               same reference, disagreeing amounts/dates: work queue #1
  Unmatched_Internal   booked but not settled (timing vs. genuine)
  Unmatched_Processor  settled but not booked
  Duplicates           repeated lines inside a file
  Rejected             rows that could not be parsed, with the reason
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .matcher import ReconResult
from .models import MatchGroup, Status

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="1F3A5F")
MONEY = '#,##0.00;[Red]-#,##0.00'
THIN = Border(top=Side(style="thin", color="999999"))

STATUS_FILL = {
    Status.MATCHED_EXACT: "E8F5E9", Status.MATCHED_TOLERANCE: "F1F8E9", Status.MATCHED_FUZZY: "FFFDE7",
    Status.MATCHED_MANY_TO_ONE: "E3F2FD", Status.MATCHED_ONE_TO_MANY: "E3F2FD", Status.BREAK: "FFEBEE",
    Status.UNMATCHED_INTERNAL: "FFF3E0", Status.UNMATCHED_PROCESSOR: "FFF3E0", Status.DUPLICATE: "F3E5F5",
}

DETAIL_COLUMNS = ["Match ID", "Status", "Rule", "Side", "Record ID", "Date", "Amount", "Currency",
                  "Reference", "Counterparty", "Group internal", "Group processor", "Difference",
                  "Days apart", "Score", "Note"]


def _detail_rows(groups: list[MatchGroup]):
    for g in groups:
        for side, recs in (("internal", g.internal), ("processor", g.processor)):
            for r in recs:
                yield [g.match_id, g.status.value, g.rule, side, r.record_id, r.txn_date, r.amount,
                       r.currency, r.reference, r.counterparty, g.internal_cents / 100,
                       g.processor_cents / 100, g.difference_cents / 100, g.date_diff_days, g.score, g.note]


def _write_table(ws, header: list[str], rows, money_cols=(), widths=None, fills=None) -> None:
    """Write a styled table. Each cell is styled as it is created (linear time even for 100k rows)."""
    money_idx = {header.index(c) for c in money_cols}
    date_idx = header.index("Date") if "Date" in header else None
    for col, name in enumerate(header, start=1):
        cell = ws.cell(1, col, name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    fill_cache: dict[str, PatternFill] = {}
    r = 1
    for r, row in enumerate(rows, start=2):
        color = fills(row) if fills else None
        fill = fill_cache.setdefault(color, PatternFill("solid", fgColor=color)) if color else None
        for c, value in enumerate(row):
            cell = ws.cell(r, c + 1, value)
            cell.font = BODY_FONT
            if fill:
                cell.fill = fill
            if c in money_idx:
                cell.number_format = MONEY
            elif c == date_idx:
                cell.number_format = "yyyy-mm-dd"
    for idx, name in enumerate(header, start=1):
        width = (widths or {}).get(name, max(12, min(45, len(name) + 4)))
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    if r > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{r}"


def _write_detail(ws, groups):
    status_of = {s.value: s for s in Status}
    _write_table(ws, DETAIL_COLUMNS, list(_detail_rows(groups)),
                 money_cols=("Amount", "Group internal", "Group processor", "Difference"),
                 widths={"Counterparty": 26, "Reference": 16, "Note": 70, "Status": 22, "Rule": 22},
                 fills=lambda row: STATUS_FILL.get(status_of[row[1]]))


def _write_summary(ws, result: ReconResult) -> None:
    s = result.summary()
    ws["A1"] = "Reconciliation summary"
    ws["A1"].font = TITLE
    ws["A2"] = f"Cutoff (as of): {s['as_of']}"
    ws["A2"].font = BODY_FONT

    ws["A4"], ws["B4"], ws["C4"] = "Headline", "Internal ledger", "Processor file"
    for c in ("A4", "B4", "C4"):
        ws[c].fill, ws[c].font = HEADER_FILL, HEADER_FONT
    rows = [
        ("Records loaded", s["internal_records"], s["processor_records"]),
        ("Total amount", s["internal_total_cents"] / 100, s["processor_total_cents"] / 100),
        ("Records matched (%)", s["internal_match_rate"], s["processor_match_rate"]),
    ]
    for i, (label, a, b) in enumerate(rows, start=5):
        ws.cell(i, 1, label).font = BODY_FONT
        ws.cell(i, 2, a).font = BODY_FONT
        ws.cell(i, 3, b).font = BODY_FONT
    for c in ("B6", "C6"):
        ws[c].number_format = MONEY
    for c in ("B7", "C7"):
        ws[c].number_format = "0.0%"
    ws.cell(8, 1, "Rows rejected at load").font = BODY_FONT
    ws.cell(8, 2, s["rejected_rows"]).font = BODY_FONT

    # Status breakdown
    r = 10
    for col, name in enumerate(("Status", "Groups", "Internal amount", "Processor amount", "Difference"), start=1):
        cell = ws.cell(r, col, name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    first = r + 1
    for status in Status:
        groups = result.by_status(status)
        r += 1
        ws.cell(r, 1, status.value).font = BODY_FONT
        ws.cell(r, 2, len(groups)).font = BODY_FONT
        ws.cell(r, 3, sum(g.internal_cents for g in groups) / 100).number_format = MONEY
        ws.cell(r, 4, sum(g.processor_cents for g in groups) / 100).number_format = MONEY
        ws.cell(r, 5, f"=C{r}-D{r}").number_format = MONEY
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
        for col in (3, 4, 5):
            ws.cell(r, col).font = BODY_FONT
    last = r
    r += 1
    ws.cell(r, 1, "Total").font = BOLD
    for col in (2, 3, 4, 5):
        letter = get_column_letter(col)
        cell = ws.cell(r, col, f"=SUM({letter}{first}:{letter}{last})")
        cell.font, cell.border = BOLD, THIN
        if col > 2:
            cell.number_format = MONEY
    status_total_row = r

    # Proof: the gap between the files is fully explained by the categories above.
    r += 2
    ws.cell(r, 1, "Reconciliation proof").font = TITLE
    r += 1
    ws.cell(r, 1, "Internal ledger total").font = BODY_FONT
    ws.cell(r, 2, "=B6").number_format = MONEY
    r += 1
    ws.cell(r, 1, "Less: processor settlement total").font = BODY_FONT
    ws.cell(r, 2, "=C6").number_format = MONEY
    r += 1
    ws.cell(r, 1, "Gap to explain").font = BOLD
    ws.cell(r, 2, f"=B{r-2}-B{r-1}").number_format = MONEY
    ws.cell(r, 2).font = BOLD
    gap_row = r
    r += 1
    ws.cell(r, 1, "Explained by status differences (Total column E above)").font = BODY_FONT
    ws.cell(r, 2, f"=E{status_total_row}").number_format = MONEY
    r += 1
    ws.cell(r, 1, "Unexplained (must be 0.00)").font = BOLD
    cell = ws.cell(r, 2, f"=ROUND(B{gap_row}-B{r-1},2)")
    cell.number_format, cell.font, cell.border = MONEY, BOLD, THIN
    ws.cell(r, 3, f'=IF(B{r}=0,"TIES OUT","CHECK")').font = BOLD

    r += 2
    ws.cell(r, 1, "Note: rejected rows are excluded from all totals; see the Rejected sheet.").font = \
        Font(name=FONT, size=9, italic=True, color="666666")

    ws.column_dimensions["A"].width = 56
    for col in "BCDE":
        ws.column_dimensions[col].width = 20


def write_excel(result: ReconResult, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _write_summary(wb.active, result)
    wb.active.title = "Summary"

    matched = [g for g in result.groups if g.status.is_match]
    _write_detail(wb.create_sheet("Matched"), matched)
    _write_detail(wb.create_sheet("Breaks"), result.by_status(Status.BREAK))
    _write_detail(wb.create_sheet("Unmatched_Internal"), result.by_status(Status.UNMATCHED_INTERNAL))
    _write_detail(wb.create_sheet("Unmatched_Processor"), result.by_status(Status.UNMATCHED_PROCESSOR))
    _write_detail(wb.create_sheet("Duplicates"), result.by_status(Status.DUPLICATE))

    rej = wb.create_sheet("Rejected")
    _write_table(rej, ["Source", "Row", "Reason", "Raw data"],
                 [[x.source, x.row_number, x.reason, str(x.raw)] for x in result.rejected],
                 widths={"Reason": 50, "Raw data": 90})

    for ws in wb.worksheets:  # print each sheet landscape, one page wide
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

    wb.calculation.fullCalcOnLoad = True  # Excel recalculates the summary formulas on open
    wb.save(path)
    return path


def write_csv(result: ReconResult, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(DETAIL_COLUMNS)
        for row in _detail_rows(result.groups):
            w.writerow([x.isoformat() if hasattr(x, "isoformat") else x for x in row])
    return path


def console_summary(result: ReconResult) -> str:
    s = result.summary()
    lines = [
        f"Reconciliation as of {s['as_of']}",
        f"  Internal records : {s['internal_records']:>6}   total ${s['internal_total_cents']/100:>14,.2f}",
        f"  Processor records: {s['processor_records']:>6}   total ${s['processor_total_cents']/100:>14,.2f}",
        f"  Rejected rows    : {s['rejected_rows']:>6}",
        f"  Match rate       : internal {s['internal_match_rate']:.1%}, processor {s['processor_match_rate']:.1%}",
        "",
        "  Groups by status:",
    ]
    for status in Status:
        n = s["groups_by_status"].get(status.value, 0)
        if n:
            lines.append(f"    {status.value:<22}{n:>6}")
    lines += ["", "  Proof:"]
    for label, cents in result.proof():
        lines.append(f"    {label:<48}{cents/100:>14,.2f}")
    return "\n".join(lines)
